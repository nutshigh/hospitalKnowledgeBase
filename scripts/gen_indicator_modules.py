"""从 data/附件2-体检信息采集模板_v1.20_最新.xlsx row2 抽取模块序。

用法:
  python3 scripts/gen_indicator_modules.py            # stdout 打印 JSON
  python3 scripts/gen_indicator_modules.py --write backend/tests/fixtures/excel_row2_modules.json

注意:运行依赖 openpyxl(系统 python 已装),backend .venv 不装;产物为静态 JSON 提交入库。
"""
import argparse
import json

import openpyxl

XLSX = "data/附件2-体检信息采集模板_v1.20_最新.xlsx"
SHEET = "体检数据采集模板"


def extract_row2_modules(path: str) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    cover = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_row == 2:
            for c in range(mr.min_col, mr.max_col + 1):
                cover[c] = mr.min_col
    items = []
    for c in range(1, ws.max_column + 1):
        if cover.get(c, c) != c:
            continue
        v = ws.cell(row=2, column=c).value
        if v is None or str(v).strip() == "":
            continue
        items.append(str(v).strip())
    return items


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--write")
    args = ap.parse_args()
    modules = extract_row2_modules(XLSX)
    print(json.dumps(modules, ensure_ascii=False, indent=1))
    if args.write:
        with open(args.write, "w", encoding="utf-8") as f:
            json.dump(modules, f, ensure_ascii=False, indent=1)
            f.write("\n")
        print("written:", args.write, "count:", len(modules))


if __name__ == "__main__":
    main()
